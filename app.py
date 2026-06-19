import streamlit as st
import os
import json
from datetime import datetime
from dotenv import load_dotenv
from anthropic import Anthropic
from pypdf import PdfReader
import docx
from docx.shared import RGBColor
import io
from fpdf import FPDF

load_dotenv()

def get_api_key():
    try:
        return st.secrets["ANTHROPIC_API_KEY"]
    except Exception:
        return os.getenv("ANTHROPIC_API_KEY")

client = Anthropic(api_key=get_api_key())

QUESTION_PROMPT = """You are a strict but helpful tutor. Based on the following study material, generate exactly ONE multiple-choice quiz question that tests real understanding, not just recall of a definition. Prefer questions that require reasoning over simple "what is X" questions.

Also assign this question a short topic label (2-4 words) representing the specific sub-topic it tests.

Create 4 answer options. Exactly one must be correct. The other 3 should be plausible but clearly wrong to someone who understands the material.
{avoid_clause}
Respond with ONLY valid JSON in this exact format, nothing else, no markdown code fences:
{{
  "topic": "short topic label",
  "question": "the question text",
  "options": {{"A": "option text", "B": "option text", "C": "option text", "D": "option text"}},
  "correct_answer": "A"
}}

STUDY MATERIAL:
{notes}
"""

GRADING_PROMPT = """You are a strict but supportive tutor giving feedback on a multiple-choice answer.

STUDY MATERIAL (ground truth):
{notes}

QUESTION:
{question}

OPTIONS:
{options}

CORRECT ANSWER: {correct_answer}
STUDENT CHOSE: {student_choice}

Write a 2-3 sentence explanation. If the student got it right, briefly reinforce why that answer is correct. If they got it wrong, explain clearly why their choice was wrong and why the correct answer is right, in a supportive tone.
"""

REEXPLAIN_PROMPT = """The student got this question wrong:

QUESTION: {question}
CORRECT ANSWER: {correct_answer}
STUDENT CHOSE: {student_choice}

Using ONLY the study material below, re-explain the underlying concept in a DIFFERENT way than a textbook definition would - use an analogy or a simpler framing. Keep it to 3-4 sentences.

STUDY MATERIAL:
{notes}
"""

def extract_text_from_pdf(file_bytes):
    reader = PdfReader(io.BytesIO(file_bytes))
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text.strip()

def extract_text_from_docx(file_bytes):
    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def extract_text_from_upload(uploaded_file):
    file_bytes = uploaded_file.read()
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
        if not text:
            return None, "This PDF appears to be scanned (no extractable text). Try a text-based PDF, or paste manually."
        return text, None
    elif name.endswith(".docx"):
        return extract_text_from_docx(file_bytes), None
    elif name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore"), None
    else:
        return None, "Unsupported file type. Please upload a PDF, DOCX, or TXT file."

def parse_json_response(text):
    text = text.strip()
    if text.startswith("```"):
        text = text.split("```")[1]
        if text.startswith("json"):
            text = text[4:]
    return json.loads(text.strip())

def generate_question(notes, avoid_topics=None):
    avoid_clause = ""
    if avoid_topics:
        avoid_clause = "\nDo not repeat these already-asked questions: " + str(avoid_topics) + "\n"
    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=500,
        messages=[{"role": "user", "content": QUESTION_PROMPT.format(notes=notes, avoid_clause=avoid_clause)}]
    )
    return parse_json_response(message.content[0].text)

def get_feedback_and_quality(notes, question, options, correct_answer, student_choice):
    options_text = "\n".join(f"{k}: {v}" for k, v in options.items())
    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=300,
        messages=[{"role": "user", "content": GRADING_PROMPT.format(
            notes=notes, question=question, options=options_text,
            correct_answer=correct_answer, student_choice=student_choice
        )}]
    )
    text = message.content[0].text.strip()
    quality = 5 if student_choice == correct_answer else 1
    return text, quality

def reexplain(notes, question, correct_answer, student_choice):
    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=250,
        messages=[{"role": "user", "content": REEXPLAIN_PROMPT.format(
            notes=notes, question=question, correct_answer=correct_answer, student_choice=student_choice
        )}]
    )
    return message.content[0].text.strip()

def generate_summary_of_notes(notes):
    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=600,
        messages=[{"role": "user", "content": (
            "Summarize the following study material into clear, well-organized study notes. "
            "Use short headings and bullet points. Focus on the key concepts a student needs to "
            "remember, not minor details. Keep it concise but complete.\n\nMATERIAL:\n" + notes
        )}]
    )
    return message.content[0].text.strip()

def build_summary_docx(summary_text):
    doc = docx.Document()
    doc.add_heading("Study Summary", level=1)
    for line in summary_text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def build_summary_pdf(summary_text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Study Summary", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 11)
    for line in summary_text.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(2)
            continue
        clean_line = line.encode("latin-1", "ignore").decode("latin-1")
        if clean_line.startswith("#"):
            pdf.set_font("Helvetica", "B", 13)
            pdf.multi_cell(0, 8, clean_line.lstrip("# "))
            pdf.set_font("Helvetica", "", 11)
        else:
            pdf.multi_cell(0, 7, clean_line)
    return bytes(pdf.output())

def build_quiz_docx(cards):
    doc = docx.Document()
    doc.add_heading("Quiz Export", level=1)
    doc.add_paragraph("Generated: " + datetime.now().strftime("%Y-%m-%d %H:%M"))
    for i, c in enumerate(cards, 1):
        doc.add_heading(f"Q{i}. [{c['topic']}]", level=2)
        doc.add_paragraph(c["question"])
        for k, v in c["options"].items():
            doc.add_paragraph(f"{k}. {v}", style="List Bullet")
        p = doc.add_paragraph()
        run = p.add_run(f"Correct answer: {c['correct_answer']}")
        run.bold = True
        run.font.color.rgb = RGBColor(0, 128, 0)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def build_quiz_pdf(cards):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Quiz Export", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 8, "Generated: " + datetime.now().strftime("%Y-%m-%d %H:%M"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for i, c in enumerate(cards, 1):
        pdf.set_font("Helvetica", "B", 12)
        q_clean = c["question"].encode("latin-1", "ignore").decode("latin-1")
        topic_clean = c["topic"].encode("latin-1", "ignore").decode("latin-1")
        pdf.multi_cell(0, 7, f"Q{i}. [{topic_clean}] {q_clean}")
        pdf.set_font("Helvetica", "", 10)
        for k, v in c["options"].items():
            v_clean = v.encode("latin-1", "ignore").decode("latin-1")
            pdf.multi_cell(0, 6, f"   {k}. {v_clean}")
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(0, 128, 0)
        pdf.multi_cell(0, 6, f"   Correct answer: {c['correct_answer']}")
        pdf.set_text_color(0, 0, 0)
        pdf.ln(3)
    return bytes(pdf.output())

st.set_page_config(page_title="Study Buddy", page_icon="book", layout="centered")
st.title("Study Buddy")
st.caption("Upload notes, get a summary, get quizzed, and download everything for later.")

defaults = {
    "notes": "", "current_card": None, "mode": "setup",
    "feedback_text": "", "reexplanation": "", "student_choice": None,
    "session_queue": [], "session_index": 0, "session_results": [],
    "generated_cards": [], "summary_text": ""
}
for key, val in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = val

st.subheader("Your study material")
tab1, tab2 = st.tabs(["Upload a file", "Paste text"])

uploaded_text = None
with tab1:
    uploaded_file = st.file_uploader("Upload notes (PDF, DOCX, or TXT)", type=["pdf", "docx", "txt"])
    if uploaded_file is not None:
        extracted, error = extract_text_from_upload(uploaded_file)
        if error:
            st.error(error)
        else:
            uploaded_text = extracted
            st.success(f"Extracted {len(extracted)} characters from {uploaded_file.name}")

with tab2:
    pasted_text = st.text_area("Paste or update your notes here", value=st.session_state.notes, height=200,
                                 placeholder="Paste lecture notes, a textbook chapter, or any material...")

notes_to_use = uploaded_text if uploaded_text else pasted_text

st.divider()
st.subheader("Summarize")

if st.button("Generate summary of these notes", use_container_width=True):
    if not notes_to_use or not notes_to_use.strip():
        st.warning("Please upload a file or paste some notes first.")
    else:
        st.session_state.notes = notes_to_use
        try:
            with st.spinner("Summarizing..."):
                st.session_state.summary_text = generate_summary_of_notes(notes_to_use)
        except Exception as e:
            st.error(f"Error generating summary: {e}")

if st.session_state.summary_text:
    st.markdown(st.session_state.summary_text)
    dl_col1, dl_col2 = st.columns(2)
    with dl_col1:
        st.download_button(
            "Download summary as Word",
            data=build_summary_docx(st.session_state.summary_text),
            file_name="study_summary.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
    with dl_col2:
        st.download_button(
            "Download summary as PDF",
            data=build_summary_pdf(st.session_state.summary_text),
            file_name="study_summary.pdf",
            mime="application/pdf",
            use_container_width=True
        )

st.divider()
st.subheader("Quiz")

num_new = st.number_input("How many questions do you want?", min_value=1, max_value=20, value=5, step=1)

col1, col2 = st.columns(2)
with col1:
    if st.button("Generate quiz", use_container_width=True):
        if not notes_to_use or not notes_to_use.strip():
            st.warning("Please upload a file or paste some notes first.")
        else:
            st.session_state.notes = notes_to_use
            progress_placeholder = st.empty()
            asked = []
            cards = []
            try:
                for i in range(num_new):
                    progress_placeholder.info(f"Generating question {i+1} of {num_new}...")
                    q = generate_question(notes_to_use, asked)
                    asked.append(q["question"])
                    cards.append(q)
                progress_placeholder.empty()
                st.session_state.generated_cards = cards
                st.session_state.session_queue = cards
                st.session_state.session_index = 0
                st.session_state.session_results = []
                st.session_state.current_card = cards[0]
                st.session_state.mode = "quiz"
                st.rerun()
            except Exception as e:
                progress_placeholder.empty()
                st.error(f"Failed after generating {len(cards)} questions. Error: {e}")

with col2:
    if st.session_state.generated_cards:
        cards = st.session_state.generated_cards
        st.download_button(
            "Download quiz as Word",
            data=build_quiz_docx(cards),
            file_name="quiz.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
        st.download_button(
            "Download quiz as PDF",
            data=build_quiz_pdf(cards),
            file_name="quiz.pdf",
            mime="application/pdf",
            use_container_width=True
        )
    else:
        st.button("Download quiz as Word", disabled=True, use_container_width=True)

if st.button("Reset everything", use_container_width=True):
    for key in defaults:
        del st.session_state[key]
    st.rerun()

st.divider()

if st.session_state.mode == "quiz" and st.session_state.current_card:
    card = st.session_state.current_card
    st.caption(f"Topic: {card['topic']} | Question {st.session_state.session_index + 1} of {len(st.session_state.session_queue)}")
    st.subheader("Question")
    st.write(card["question"])

    choice = st.radio(
        "Choose an answer:",
        options=list(card["options"].keys()),
        format_func=lambda k: f"{k}. {card['options'][k]}",
        key=f"mc_choice_{st.session_state.session_index}",
        index=None
    )

    if st.button("Submit answer", use_container_width=True):
        if choice is None:
            st.warning("Please select an answer.")
        else:
            is_correct = (choice == card["correct_answer"])
            st.session_state.student_choice = choice
            st.session_state.session_results.append({
                "question": card["question"], "topic": card["topic"], "correct": is_correct
            })
            try:
                with st.spinner("Checking your answer..."):
                    feedback, quality = get_feedback_and_quality(
                        st.session_state.notes, card["question"], card["options"],
                        card["correct_answer"], choice
                    )
                    st.session_state.feedback_text = feedback
                    if not is_correct:
                        st.session_state.reexplanation = reexplain(
                            st.session_state.notes, card["question"], card["correct_answer"], choice
                        )
                        st.session_state.mode = "reteach"
                    else:
                        st.session_state.mode = "graded"
                st.rerun()
            except Exception as e:
                st.error(f"Error checking answer: {e}")

def advance_session():
    st.session_state.session_index += 1
    if st.session_state.session_index >= len(st.session_state.session_queue):
        st.session_state.mode = "summary"
    else:
        st.session_state.current_card = st.session_state.session_queue[st.session_state.session_index]
        st.session_state.mode = "quiz"
    st.rerun()

if st.session_state.mode == "graded":
    card = st.session_state.current_card
    st.success(f"Correct! You chose {st.session_state.student_choice}: {card['options'][st.session_state.student_choice]}")
    st.write(st.session_state.feedback_text)
    label = "See results" if st.session_state.session_index + 1 >= len(st.session_state.session_queue) else "Next question"
    if st.button(label, use_container_width=True):
        advance_session()

if st.session_state.mode == "reteach":
    card = st.session_state.current_card
    st.warning(f"Not quite. You chose {st.session_state.student_choice}: {card['options'][st.session_state.student_choice]}")
    st.write(f"Correct answer was {card['correct_answer']}: {card['options'][card['correct_answer']]}")
    st.write(st.session_state.feedback_text)
    st.info("Let's look at this differently:\n\n" + st.session_state.reexplanation)
    label = "See results" if st.session_state.session_index + 1 >= len(st.session_state.session_queue) else "Next question"
    if st.button(label, use_container_width=True):
        advance_session()

if st.session_state.mode == "summary":
    correct_count = sum(1 for r in st.session_state.session_results if r["correct"])
    total = len(st.session_state.session_results)
    st.subheader(f"Quiz complete! {correct_count}/{total} correct")
    st.progress(correct_count / total if total else 0)

    with st.expander("See all questions and answers"):
        for i, r in enumerate(st.session_state.session_results):
            mark = "CORRECT" if r["correct"] else "WRONG"
            st.write(f"[{mark}] Q{i+1} ({r['topic']}): {r['question']}")

    if st.button("Back to study material", use_container_width=True):
        st.session_state.mode = "setup"
        st.rerun()

st.divider()
st.caption("Built by Sara Ghassemi - Powered by Claude Anthropic")
